"""Download the Vikaspuri project workbook and turn its tabs into clean tables.

Only the data tabs are parsed (Building, Land, Other, Land Reference, Consolidated Summary,
Builder Agreement). The summary tabs are formulas over the same rows, so every total is
recomputed from the rows here instead of being read from them.
"""

import hashlib
import io
from dataclasses import dataclass

import openpyxl
import pandas as pd
import requests

EXPORT_URL = "https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"

PAYMENT_MODES = {
    "online": "Online",
    "cash": "Cash",
    "upi": "UPI",
    "online+cash": "Online+Cash",
    "unspecified": "Unspecified",
}

SECTION_NAMES = {"building": "Building construction", "land": "Land purchase", "other": "Other related"}

MILESTONES_FACT = "Construction milestones"


@dataclass
class ProjectData:
    expenses: pd.DataFrame  # section, date, payee, item, category, mode, amount, notes
    milestones: pd.DataFrame  # date, event, paid_so_far
    agreement: pd.DataFrame  # ref, item, allowance, notes, extra_cost
    agreement_summary: dict[str, str]
    facts: dict[str, str]
    owner_count: int
    version: str


def download_workbook(sheet_id: str) -> bytes:
    response = requests.get(EXPORT_URL.format(sheet_id=sheet_id), timeout=30)
    response.raise_for_status()
    return response.content


def parse_workbook(content: bytes) -> ProjectData:
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)

    building = _read_table(wb["Building Expenses"], "S.No.")
    land = _read_table(wb["Land Expenses"], "S.No.")
    other = _read_table(wb["Other Related Expenses"], "S.No.")

    rows = []
    for r in building:
        rows.append(_expense("building", r, "Expense Details", payee=r.get("Person / Payee"), notes=r.get("Notes")))
    for r in land:
        rows.append(_expense("land", r, "Expense Item", notes=r.get("Comments")))
    for r in other:
        rows.append(_expense("other", r, "Expense Item", category=r.get("Category"), notes=r.get("Comments")))
    expenses = pd.DataFrame(rows)
    expenses = expenses[expenses["amount"] > 0].reset_index(drop=True)
    expenses["date"] = pd.to_datetime(expenses["date"], errors="coerce")

    milestones = pd.DataFrame(
        [
            {"date": r.get("Date"), "event": r.get("Milestones"), "paid_so_far": _number(r.get("Milestone Amount"))}
            for r in building
            if r.get("Milestones")
        ],
        columns=["date", "event", "paid_so_far"],
    )
    milestones["date"] = pd.to_datetime(milestones["date"], errors="coerce")

    facts, owner_count = _read_facts(wb)
    if not milestones.empty:
        facts[MILESTONES_FACT] = "; ".join(milestones["event"])
    agreement, agreement_summary = _read_agreement(wb["Builder Agreement Items"])

    return ProjectData(
        expenses=expenses,
        milestones=milestones,
        agreement=agreement,
        agreement_summary=agreement_summary,
        facts=facts,
        owner_count=owner_count,
        version=hashlib.sha256(content).hexdigest()[:12],
    )


def inr(amount: float) -> str:
    """Format rupees with Indian digit grouping, e.g. 9400000 -> ₹94,00,000."""
    rupees = round(amount)
    digits = str(abs(rupees))
    head, tail = digits[:-3], digits[-3:]
    groups = []
    while len(head) > 2:
        groups.insert(0, head[-2:])
        head = head[:-2]
    if head:
        groups.insert(0, head)
    return ("-" if rupees < 0 else "") + "₹" + ",".join(groups + [tail])


def inr_short(amount: float) -> str:
    """Compact lakh/crore form, e.g. 9400000 -> ₹94.0 L."""
    if abs(amount) >= 1e7:
        return f"₹{amount / 1e7:.2f} Cr"
    if abs(amount) >= 1e5:
        return f"₹{amount / 1e5:.1f} L"
    return inr(amount)


def _rows(ws):
    for row in ws.iter_rows(values_only=True):
        yield [v.strip() or None if isinstance(v, str) else v for v in row]


def _read_table(ws, first_header: str, stop_prefix: str | None = None) -> list[dict]:
    """Rows below the header row whose first cell is `first_header`, up to the 'Total' row."""
    header = None
    records = []
    for row in _rows(ws):
        if header is None:
            if row and row[0] == first_header:
                header = row
            continue
        if any(isinstance(v, str) and v.endswith("Total") for v in row):
            break
        if stop_prefix and isinstance(row[0], str) and row[0].startswith(stop_prefix):
            break
        if all(v is None for v in row):
            continue
        records.append({h: v for h, v in zip(header, row) if h is not None})
    return records


def _expense(section: str, r: dict, item_key: str, payee=None, category=None, notes=None) -> dict:
    mode = r.get("Payment Mode")
    return {
        "section": section,
        "date": r.get("Date"),
        "payee": payee,
        "item": r.get(item_key),
        "category": category,
        "mode": PAYMENT_MODES.get(str(mode).lower(), str(mode)) if mode else "Unspecified",
        "amount": _number(r.get("Amount (INR)")) or 0.0,
        "notes": None if notes is None else str(notes),
    }


def _number(value) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        digits = "".join(ch for ch in value if ch.isdigit() or ch == ".")
        return float(digits) if digits else None
    return None


def _read_facts(wb) -> tuple[dict[str, str], int]:
    facts: dict[str, str] = {}

    summary = list(_rows(wb["Consolidated Summary"]))
    # Floor table: "Floor | Owners | Children | ..."; every titled column becomes a per-floor fact
    # plus one combined fact, so new columns added to the sheet are picked up automatically.
    floor_columns: list[tuple[int, str]] = []
    combined: dict[str, list[str]] = {}
    for i, row in enumerate(summary):
        first = row[0] if row else None
        if first == "Floor":
            floor_columns = [(j, str(h).strip().lower()) for j, h in enumerate(row) if j > 0 and h]
            continue
        if floor_columns and isinstance(first, str) and first.endswith("Floor"):
            for j, header in floor_columns:
                value = str(row[j]).strip().rstrip(".") if j < len(row) and row[j] else None
                if value:
                    facts[f"{first} {header}"] = value
                    combined.setdefault(header, []).append(f"{first}: {value}")
            continue
        if floor_columns:
            for header, lines in combined.items():
                facts[f"{header.capitalize()} (all floors)"] = "\n".join(lines)
            floor_columns = []
        if isinstance(first, str) and first.endswith("details:"):
            detail = next((r[0] for r in summary[i + 1 :] if r and r[0] is not None), None)
            if isinstance(detail, str) and not detail.startswith("Enter any") and not detail.endswith("details:"):
                facts[first.rstrip(":").strip()] = detail

    owner_count = 4
    reference = wb["Land Reference Details"]
    for row in _rows(reference):
        if row and row[0] == "Seller":
            facts["Land seller"] = row[1]
            if len(row) > 4 and row[3] == "Banker":
                facts["Bank providing the loan"] = row[4]
    for r in _read_table(reference, "Reference Item"):
        label, value = r.get("Reference Item"), r.get("Value")
        if label is None or value is None:
            continue
        if label.startswith("Total persons"):
            owner_count = int(value)
        facts[label] = _format_reference(label, value)
    return facts, owner_count


def _format_reference(label: str, value) -> str:
    if not isinstance(value, (int, float)):
        return str(value)
    lowered = label.lower()
    if "%" in label:
        return f"{value:g}%"
    if any(word in lowered for word in ("amount", "cost", "loan")):
        return f"{inr(value)} ({inr_short(value)})"
    return f"{value:g}"


def _read_agreement(ws) -> tuple[pd.DataFrame, dict[str, str]]:
    summary: dict[str, str] = {}
    for row in _rows(ws):
        if row and row[0] == "Ref":
            break
        if len(row) > 4 and isinstance(row[3], str) and row[4] is not None:
            summary[row[3]] = str(row[4])
    for row in _rows(ws):
        if row and isinstance(row[0], str) and row[0].startswith("Important"):
            summary["Note"] = row[0]

    records = []
    for r in _read_table(ws, "Ref", stop_prefix="Important"):
        ref, item, allowance, notes, extra = (list(r.values()) + [None] * 5)[:5]
        if isinstance(ref, float) and ref.is_integer():
            ref = int(ref)
        records.append(
            {"ref": str(ref), "item": item, "allowance": allowance, "notes": notes, "extra_cost": _number(extra)}
        )
    return pd.DataFrame(records), summary
